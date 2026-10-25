"""NC-4: persisted, scoped NovaCase results and exports."""
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from pypdf import PdfReader
from docx import Document

from flask import g, request

from app import create_app
from app.config import Config
from app.models.task import TaskManager, TaskStatus
from app.services.case_corpus import CaseCorpusRepository, DocumentIngestionService
from app.services.novacourt_pipeline_service import NovaCourtPipelineService
from app.services.runtime_security import CaseOwnershipRepository
from app.services.task_repository import TaskRepository
from app.services.case_pdf import render_case_pdf
from app.services.case_report_content import report_blocks


class Embeddings:
    def generate_embeddings(self, texts):
        return [[0.1, 0.2, 0.3] for _ in texts]


class NovaCaseNC4Tests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        task_db, corpus_db = root / 'tasks.sqlite3', root / 'corpus.sqlite3'

        class TestConfig(Config):
            TESTING = True
            RATE_LIMIT_ENABLED = False
            TASK_PERSISTENCE_ENABLED = True
            TASK_DB_PATH = str(task_db)
            CASE_CORPUS_DB_PATH = str(corpus_db)

        self.app = create_app(TestConfig)

        @self.app.before_request
        def test_principal():
            g.authenticated_user_id = request.headers.get('X-Test-Owner', 'owner-a')

        self.client = self.app.test_client()
        self.owners = CaseOwnershipRepository(corpus_db)
        self.assertTrue(self.owners.owned('CASE-A', 'owner-a', create=True))
        self.assertTrue(self.owners.owned('CASE-B', 'owner-b', create=True))
        self.tasks = self.app.extensions['task_manager']
        self.task_id = self.tasks.create_task('legal_analysis',
            {'owner_id': 'owner-a', 'case_id': 'CASE-A', 'tool': 'case',
             'knowledge_version': 'v2', 'document_ids': []})
        self.result = {'success': True, 'case_id': 'CASE-A', 'case': 'Caso A',
                       'report': 'Informe canónico', 'citations': [], 'sources': []}

    def tearDown(self):
        self.temp.cleanup()

    def request_export(self, case_id='CASE-A', task_id=None, owner='owner-a', **extra):
        return self.client.post('/api/export/prepare', json={
            'case_id': case_id, 'task_id': task_id or self.task_id, 'tool': 'case', **extra},
            headers={'X-Test-Owner': owner})

    def test_novacourt_exports_use_persisted_scoped_result_and_share_canonical_content(self):
        task_id = self.tasks.create_task('legal_analysis',
            {'owner_id':'owner-a', 'case_id':'CASE-A', 'tool':'court', 'document_ids':[]})
        private = {'source_id':'SRC-private-doc', 'source_scope':'private_case', 'case_id':'CASE-A',
                   'display_title':'Contrato CAS aportado', 'page_start':1}
        legal = {'source_id':'SRC-public-unit', 'source_scope':'public_legal',
                 'display_title':'STC Exp. 02047-2025-PA/TC', 'court':'Tribunal Constitucional',
                 'unit_type':'foundation', 'unit_number':'11', 'page_start':5, 'page_end':6}
        canonical = {**self.result, 'simulation':{'status':'ready'},
            'court_report_document':{'document_type':'court_simulation', 'case_id':'CASE-A',
                'case_title':'Municipalidad Distrital de Santa Esperanza',
                'sections':[{'title':'Posición de la parte demandante', 'content':'## Tesis principal\nLa no renovación se cuestiona. [1]'},
                            {'title':'Decisión simulada', 'content':'La conclusión depende de acreditar los hechos. [4]'}],
                'citations':[{'citation_id':'CIT-1','label':'[1]','source_id':private['source_id']},
                             {'citation_id':'CIT-4','label':'[4]','source_id':legal['source_id']}],
                'sources':[private, legal]}}
        self.tasks.update_task(task_id, status=TaskStatus.PROCESSING)
        self.tasks.complete_task(task_id, canonical)

        for fmt, mime in (('pdf','application/pdf'),
                          ('docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')):
            with self.subTest(format=fmt):
                payload = {'case_id':'CASE-A','task_id':task_id,'tool':'court'}
                response = self.client.post(f'/api/export/court/{fmt}', json=payload, headers={'X-Test-Owner':'owner-a'})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.mimetype, mime)
                self.assertIn('MYKE_NovaCourt_Municipalidad_Distrital_de_Santa_Esperanza',
                              response.headers['Content-Disposition'])
                self.assertGreater(len(response.data), 1000)
                if fmt == 'pdf':
                    self.assertTrue(response.data.startswith(b'%PDF'))
                    reader = PdfReader(io.BytesIO(response.data))
                    text = '\n'.join(page.extract_text() or '' for page in reader.pages)
                    self.assertAlmostEqual(float(reader.pages[0].mediabox.width), 595.3, delta=1)
                    self.assertIn('Generado por MYKE · NovaCourt', text)
                else:
                    self.assertTrue(zipfile.is_zipfile(io.BytesIO(response.data)))
                    self.assertTrue(response.data.startswith(b'PK'))
                    document = Document(io.BytesIO(response.data))
                    self.assertAlmostEqual(document.sections[0].page_width.cm, 21, delta=.1)
                    self.assertAlmostEqual(document.sections[0].left_margin.cm, 2.5, delta=.1)
                    text = '\n'.join(p.text for p in document.paragraphs)
                    text += '\n' + '\n'.join(p.text for p in document.sections[0].footer.paragraphs)
                self.assertIn('La no renovación se cuestiona. [1]', text)
                self.assertIn('La conclusión depende de acreditar los hechos. [4]', text)
                self.assertIn('Contrato CAS aportado', text)
                self.assertIn('STC Exp. 02047-2025-PA/TC', text)
                self.assertNotIn('CASE-A', text)
                self.assertNotIn('SRC-private-doc', text)
                forged = self.client.post(f'/api/export/court/{fmt}', json={**payload,'report':'TEXTO MANIPULADO'},
                                          headers={'X-Test-Owner':'owner-a'})
                self.assertEqual(forged.status_code, 400)
                self.assertNotIn('TEXTO MANIPULADO', forged.get_data(as_text=True))
                self.assertEqual(self.client.post(f'/api/export/court/{fmt}', json=payload,
                    headers={'X-Test-Owner':'owner-b'}).status_code, 404)
                self.assertEqual(self.client.post(f'/api/export/court/{fmt}', json={**payload,'case_id':'CASE-B'},
                    headers={'X-Test-Owner':'owner-a'}).status_code, 404)
                self.assertEqual(self.client.post(f'/api/export/court/{fmt}', json={**payload,'task_id':self.task_id},
                    headers={'X-Test-Owner':'owner-a'}).status_code, 404)

        pending = self.tasks.create_task('legal_analysis',
            {'owner_id':'owner-a','case_id':'CASE-A','tool':'court','document_ids':[]})
        self.assertEqual(self.client.post('/api/export/court/pdf', json={'case_id':'CASE-A','task_id':pending,'tool':'court'},
            headers={'X-Test-Owner':'owner-a'}).status_code, 409)
        self.app.extensions['task_repository'] = TaskRepository(self.app.config['TASK_DB_PATH'])
        persisted = self.client.post('/api/export/court/pdf', json={'case_id':'CASE-A','task_id':task_id,'tool':'court'},
                                     headers={'X-Test-Owner':'owner-a'})
        self.assertEqual(persisted.status_code, 200)

    def test_scoped_pdf_uses_persisted_report_and_single_footer_mark(self):
        payload = {'case_id':'CASE-A', 'task_id':self.task_id, 'tool':'case'}
        def pdf(owner='owner-a', data=None):
            return self.client.post('/api/export/pdf', json=data or payload,
                                    headers={'X-Test-Owner':owner})
        self.assertEqual(pdf().status_code, 409)
        self.tasks.update_task(self.task_id, status=TaskStatus.PROCESSING)
        self.tasks.complete_task(self.task_id, {**self.result,
            'report':'## Objeto y alcance\nHecho documentado para revisión.\n\n## Conclusión\nRevisar la prueba.',
            'report_document':{'case_id':'CASE-A', 'title':'Informe jurídico',
                'sections':[{'title':'Objeto y alcance', 'content':'Hecho documentado para revisión.'},
                            {'title':'Conclusión', 'content':'Revisar la prueba [1].'}],
                'citations':[{'label':'[1]', 'source_id':'SRC-1'}],
                'sources':[{'source_id':'SRC-1', 'title':'Documento de prueba', 'page_start':2}]}})
        response = pdf()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/pdf')
        reader = PdfReader(io.BytesIO(response.data))
        content = '\n'.join(page.extract_text() or '' for page in reader.pages)
        self.assertIn('Hecho documentado', content)
        self.assertIn('Conclusión', content)
        self.assertIn('Documento de prueba', content)
        self.assertEqual(content.count('Generado por MYKE'), 1)
        self.assertNotIn('CASE-A', content)
        self.assertEqual(pdf(owner='owner-b').status_code, 404)
        self.assertEqual(pdf(data={**payload, 'case_id':'CASE-B'}).status_code, 404)
        self.assertEqual(pdf(data={**payload, 'tool':'court'}).status_code, 400)
        self.assertEqual(pdf(data={**payload, 'report':'Texto impuesto'}).status_code, 400)

    def test_pdf_footer_once_on_each_page(self):
        document = {'case_id': 'CASE-A', 'sections': [
            {'title': 'Análisis jurídico', 'content': ('Un párrafo sintético para revisar la paginación.\n' * 180)}]}
        data = render_case_pdf({'success': True, 'case_id': 'CASE-A',
                                'report': 'Informe sintético', 'report_document': document})
        reader = PdfReader(io.BytesIO(data))
        self.assertGreaterEqual(len(reader.pages), 2)
        for page in reader.pages:
            self.assertEqual((page.extract_text() or '').count('Generado por MYKE'), 1)

    def test_docx_and_pdf_share_persisted_report_and_scope(self):
        payload = {'case_id': 'CASE-A', 'task_id': self.task_id, 'tool': 'case'}
        report = ('# Informe jurídico de análisis\n'
                  '## Objeto y alcance\nSe examina el contrato aportado [1].\n'
                  '## Conclusiones\n1. La alegación requiere corroboración.\n'
                  '## Fuentes\nReferencia duplicada que no debe presentarse.')
        source = {'source_id': 'SRC-1', 'case_id': 'CASE-A', 'source_scope': 'private_case',
                  'display_title': 'Contrato aportado', 'page_start': 1}
        self.tasks.update_task(self.task_id, status=TaskStatus.PROCESSING)
        self.tasks.complete_task(self.task_id, {**self.result, 'report': report,
            'report_document': {'case_id': 'CASE-A', 'title': 'Informe jurídico de análisis',
                'sections': [{'title': 'Objeto y alcance', 'content': 'Se examina el contrato aportado [1].'},
                             {'title': 'Conclusiones', 'content': '1. La alegación requiere corroboración.'},
                             {'title': 'Fuentes', 'content': 'Referencia duplicada que no debe presentarse.'}],
                'citations': [{'citation_id': 'CIT-1', 'label': '[1]', 'source_id': 'SRC-1'}],
                'sources': [source]}})
        for fmt in ('pdf', 'docx'):
            with self.subTest(fmt=fmt):
                def export(data=payload, owner='owner-a'):
                    return self.client.post(f'/api/export/{fmt}', json=data,
                                            headers={'X-Test-Owner': owner})
                response = export()
                self.assertEqual(response.status_code, 200)
                self.assertGreater(len(response.data), 1000)
                if fmt == 'pdf':
                    self.assertEqual(response.mimetype, 'application/pdf')
                    page = PdfReader(io.BytesIO(response.data)).pages[0]
                    text = page.extract_text()
                    self.assertAlmostEqual(float(page.mediabox.width), 595.3, delta=1)
                else:
                    self.assertEqual(response.mimetype,
                        'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
                    self.assertTrue(response.data.startswith(b'PK'))
                    document = Document(io.BytesIO(response.data))
                    self.assertAlmostEqual(document.sections[0].page_width.cm, 21, delta=.1)
                    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
                    text += '\n' + '\n'.join(paragraph.text for paragraph in document.sections[0].footer.paragraphs)
                self.assertIn('Se examina el contrato aportado [1]', text)
                self.assertIn('Contrato aportado', text)
                self.assertIn('Generado por MYKE', text)
                self.assertNotIn('Referencia duplicada', text)
                self.assertNotIn('SRC-1', text)
                self.assertEqual(export(owner='owner-b').status_code, 404)
                self.assertEqual(export(data={**payload, 'case_id': 'CASE-B'}).status_code, 404)
                self.assertEqual(export(data={**payload, 'tool': 'court'}).status_code, 400)
                self.assertEqual(export(data={**payload, 'report': 'Texto impuesto'}).status_code, 400)

    def test_narrative_sections_keep_prose_while_action_sections_keep_lists(self):
        self.assertEqual(report_blocks('- Párrafo jurídico.', 'Análisis jurídico'),
                         [('paragraph', 'Párrafo jurídico.')])
        self.assertEqual(report_blocks('- Acopiar el contrato.', 'Estrategia y actuaciones recomendadas'),
                         [('list', 'Acopiar el contrato.')])

    def test_restart_scoped_export_and_isolation(self):
        waiting = self.request_export()
        self.assertEqual(waiting.status_code, 409)
        self.assertEqual(waiting.json['error']['code'], 'TASK_NOT_READY')
        self.tasks.update_task(self.task_id, status=TaskStatus.PROCESSING,
                               progress_detail={'current_stage': 'report', 'partial_result': self.result})
        self.tasks.complete_task(self.task_id, self.result)

        repository = TaskRepository(self.app.config['TASK_DB_PATH'])
        restored_manager = TaskManager(repository)
        self.app.extensions['task_repository'] = repository
        self.app.extensions['task_manager'] = restored_manager
        self.app.extensions['legal_pipeline'] = NovaCourtPipelineService(None, None, None, restored_manager)
        restored = restored_manager.get_task(self.task_id)
        self.assertEqual(restored.status, TaskStatus.COMPLETED)
        self.assertEqual(restored.result, self.result)
        self.assertEqual(restored.metadata['knowledge_version'], 'v2')
        self.assertNotIn('partial_result', restored.progress_detail)
        self.assertEqual(repository.load_scoped('owner-a', 'CASE-A', self.task_id)['result'], self.result)
        self.assertIsNone(repository.load_scoped('owner-b', 'CASE-A', self.task_id))
        self.assertIsNone(repository.load_scoped('owner-a', 'CASE-B', self.task_id))

        status = self.client.get(f'/api/tasks/{self.task_id}?case_id=CASE-A')
        self.assertEqual(status.status_code, 200)
        self.assertEqual(status.json['final_result'], self.result)
        exported = self.request_export()
        self.assertEqual(exported.status_code, 200)
        self.assertIn('Informe canónico', str(exported.json['document']))
        self.assertEqual(self.request_export(case_id='CASE-B').status_code, 404)
        self.assertEqual(self.request_export(owner='owner-b').status_code, 404)
        self.assertEqual(self.client.get(f'/api/tasks/{self.task_id}?case_id=CASE-A',
            headers={'X-Test-Owner': 'owner-b'}).status_code, 404)

        other_id = restored_manager.create_task('legal_analysis',
            {'owner_id': 'owner-a', 'case_id': 'CASE-A', 'tool': 'case', 'knowledge_version': 'v2'})
        restored_manager.update_task(other_id, status=TaskStatus.PROCESSING)
        restored_manager.complete_task(other_id, {**self.result, 'report': 'Segundo informe'})
        self.assertIn('Informe canónico', str(self.request_export().json['document']))
        self.assertIn('Segundo informe', str(self.request_export(task_id=other_id).json['document']))
        self.assertEqual(self.request_export(task_id='00000000-0000-4000-8000-000000000000').status_code, 404)
        forged = self.request_export(result={'success': True, 'case_id': 'CASE-A', 'report': 'FALSIFICADO'})
        self.assertEqual(forged.status_code, 400)
        self.assertNotIn('FALSIFICADO', forged.get_data(as_text=True))
        malformed_id = restored_manager.create_task('legal_analysis',
            {'owner_id': 'owner-a', 'case_id': 'CASE-A', 'tool': 'case'})
        restored_manager.update_task(malformed_id, status=TaskStatus.PROCESSING)
        restored_manager.complete_task(malformed_id, {**self.result,
            'metadata': {'case_id': 'CASE-B', 'internal': 'SECRET_DSN'}})
        invalid = self.request_export(task_id=malformed_id)
        self.assertEqual(invalid.status_code, 409)
        self.assertEqual(invalid.json['error']['code'], 'TASK_NOT_READY')
        self.assertNotIn('SECRET_DSN', invalid.get_data(as_text=True))

    def test_private_document_persists_with_case_and_upload_rejects_other_owner(self):
        repository = CaseCorpusRepository(self.app.config['CASE_CORPUS_DB_PATH'])
        ingest = DocumentIngestionService(repository, Embeddings(), max_file_size=100000,
            max_total_size=200000, max_files=20, max_uncompressed=100000,
            chunk_tokens=20, chunk_overlap=2, batch_size=4, ocr_min_chars=10)
        manifest, duplicate = ingest.ingest('CASE-A', '../evidence.txt',
            io.BytesIO(b'Private alpha evidence about the case.'))
        self.assertFalse(duplicate)
        recreated = CaseCorpusRepository(self.app.config['CASE_CORPUS_DB_PATH'])
        self.assertEqual(recreated.list_documents('CASE-A')[0]['document_id'], manifest['document_id'])
        self.assertTrue(recreated.retrieve('CASE-A', [manifest['document_id']]))
        self.assertEqual(recreated.list_documents('CASE-B'), [])
        self.app.extensions.pop('case_corpus_repository', None)
        with self.app.app_context():
            from app.services.case_service import CaseService
            with patch('app.services.case_service.SearchService'):
                service = CaseService()
            self.assertEqual(service.case_corpus_repository.list_documents('CASE-A')[0]['document_id'],
                             manifest['document_id'])
        response = self.client.post('/api/cases/CASE-A/documents',
            data={'file': (io.BytesIO(b'other evidence'), 'other.txt')},
            headers={'X-Test-Owner': 'owner-b'})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.owners.owned('CASE-A', 'owner-b', create=False), False)


if __name__ == '__main__':
    unittest.main()
