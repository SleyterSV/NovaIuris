import io
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import fitz
from docx import Document

from app.services.case_corpus import (CaseContextService, CaseCorpusRepository,
    DocumentError, DocumentIngestionService)


class FakeEmbeddings:
    def __init__(self):
        self.batches = []

    @staticmethod
    def vector(text):
        words = set(text.lower().split())
        return [float(item in words) for item in ("alpha", "beta", "evidence")]

    def generate_embeddings(self, texts):
        self.batches.append(list(texts))
        return [self.vector(text) for text in texts]

    def generate_embedding(self, text, cancellation_token=None):
        return self.vector(text)


class CorpusTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.repo = CaseCorpusRepository(str(Path(self.tmp.name) / "corpus.sqlite3"))
        self.embeddings = FakeEmbeddings()
        self.service = DocumentIngestionService(self.repo, self.embeddings,
            max_file_size=2_000_000, max_total_size=4_000_000, max_files=20,
            max_uncompressed=2_000_000, chunk_tokens=16, chunk_overlap=3,
            batch_size=2, ocr_min_chars=20)

    def tearDown(self):
        self.tmp.cleanup()

    @staticmethod
    def txt(text):
        return io.BytesIO(text.encode("utf-8"))

    def test_txt_manifest_chunks_batch_hash_and_dedupe(self):
        text = "FUNDAMENTOS DE HECHO\n\nAlpha evidence confirms the first relevant event.\n\nPETITORIO\n\nBeta asks the court for relief."
        manifest, duplicate = self.service.ingest("CASE-A", "../record.txt", self.txt(text))
        self.assertFalse(duplicate)
        self.assertEqual(manifest["filename"], "record.txt")
        self.assertEqual(manifest["status"], "ready")
        self.assertEqual(manifest["mime_type"], "text/plain")
        chunks = self.repo.retrieve("CASE-A", [manifest["document_id"]])
        self.assertTrue(chunks)
        self.assertTrue(all(c["case_id"] == "CASE-A" for c in chunks))
        self.assertTrue(all(c["source_reference"]["document_id"] == manifest["document_id"] for c in chunks))
        self.assertTrue(all(c["page_start"] is None for c in chunks))
        self.assertIn("PETITORIO", {c["section"] for c in chunks})
        self.assertTrue(any("alpha" in " ".join(batch).lower() for batch in self.embeddings.batches))
        batches_before = len(self.embeddings.batches)
        again, duplicate = self.service.ingest("CASE-A", "renamed.txt", self.txt(text))
        self.assertTrue(duplicate)
        self.assertEqual(again["document_id"], manifest["document_id"])
        self.assertEqual(len(self.embeddings.batches), batches_before)
        other_case, duplicate = self.service.ingest("CASE-B", "record.txt", self.txt(text))
        self.assertFalse(duplicate)
        self.assertNotEqual(other_case["document_id"], manifest["document_id"])
        context = CaseContextService(self.repo, self.embeddings).relevant_context(
            "CASE-A", "alpha evidence", document_ids=[manifest["document_id"]])
        self.assertTrue(context)
        self.assertTrue(all(item["document_id"] == manifest["document_id"] for item in context))
        with self.assertRaises(DocumentError):
            self.repo.retrieve("CASE-B", [manifest["document_id"]])
        calls = []
        class CountingEmbedding(FakeEmbeddings):
            def generate_embedding(self, text, cancellation_token=None):
                calls.append(text)
                return super().generate_embedding(text, cancellation_token)
        with self.assertRaises(DocumentError):
            CaseContextService(self.repo, CountingEmbedding()).relevant_context(
                "CASE-B", "do not embed for another case", document_ids=[manifest["document_id"]])
        self.assertEqual(calls, [])

    def test_missing_case_is_never_an_unscoped_search(self):
        with self.assertRaises(DocumentError):
            self.repo.retrieve("")
        with self.assertRaises(DocumentError):
            CaseContextService(self.repo, self.embeddings).relevant_context("", "alpha")

    def test_storage_foreign_key_rejects_chunk_case_mismatch(self):
        manifest, _ = self.service.ingest("CASE-A", "brief.txt", self.txt("Alpha evidence belongs to CASE A."))
        chunk = self.repo.retrieve("CASE-A", [manifest["document_id"]])[0]
        import sqlite3
        with self.assertRaises(sqlite3.IntegrityError):
            with self.repo._connect() as db:
                db.execute("INSERT INTO case_chunks(chunk_id,case_id,document_id,chunk_hash,page_start,page_end,section,source_json,text,embedding_json) VALUES(?,?,?,?,?,?,?,?,?,?)",
                    ("BAD-CHUNK", "CASE-B", manifest["document_id"], "bad-hash", None, None, None,
                     "{}", "foreign chunk", "[]"))

    def test_docx_preserves_order_headings_tables_and_no_fake_page(self):
        path = Path(self.tmp.name) / "input.docx"
        document = Document()
        document.add_heading("FUNDAMENTOS DE DERECHO", level=1)
        document.add_paragraph("Alpha legal basis")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Evidence", "Annex A"
        document.save(path)
        with path.open("rb") as source:
            manifest, _ = self.service.ingest("CASE-A", "brief.docx", source)
        chunks = self.repo.retrieve("CASE-A", [manifest["document_id"]])
        self.assertTrue(all(c["page_start"] is None for c in chunks))
        combined = " ".join(c["text"] for c in chunks)
        self.assertLess(combined.index("Alpha legal basis"), combined.index("Evidence | Annex A"))
        self.assertEqual(chunks[0]["section"], "FUNDAMENTOS DE DERECHO")

    def test_pdf_page_references_and_ocr_required(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Alpha evidence from the first page with enough useful text to exceed threshold.")
        pdf.new_page()
        data = pdf.tobytes()
        pdf.close()
        manifest, _ = self.service.ingest("CASE-A", "brief.pdf", io.BytesIO(data))
        self.assertEqual(manifest["page_count"], 2)
        self.assertTrue(manifest["ocr_required"])
        self.assertEqual(manifest["warnings"][0]["page_number"], 2)
        chunks = self.repo.retrieve("CASE-A", [manifest["document_id"]])
        self.assertTrue(all(c["page_start"] == 1 for c in chunks))
        self.assertEqual(chunks[0]["source_reference"]["page_end"], 1)

    def test_replacement_is_versioned_and_old_chunks_are_not_active(self):
        first, _ = self.service.ingest("CASE-A", "brief.txt", self.txt("Alpha evidence document original."))
        second, _ = self.service.ingest("CASE-A", "brief.txt", self.txt("Beta evidence document updated."),
                                        replaces_document_id=first["document_id"])
        self.assertEqual(second["document_version"], 2)
        self.assertEqual([d["document_id"] for d in self.repo.list_documents("CASE-A")], [second["document_id"]])
        self.assertFalse(any(c["document_id"] == first["document_id"] for c in self.repo.retrieve("CASE-A")))

    def test_chunks_have_stable_hashes_and_configured_character_bound(self):
        service = DocumentIngestionService(self.repo, self.embeddings,
            max_file_size=2_000_000, max_uncompressed=2_000_000, chunk_tokens=100,
            chunk_overlap=3, batch_size=10, ocr_min_chars=20, chunk_max_chars=10)
        text = "abcdefghijklmnopqrstuvxyz0123456789"
        first = service._chunk("CASE-A", [{"text": text, "page": 3, "section": None,
                                            "paragraph_start": 1, "paragraph_end": 1}])
        second = service._chunk("CASE-A", [{"text": text, "page": 3, "section": None,
                                             "paragraph_start": 1, "paragraph_end": 1}])
        self.assertEqual("".join(chunk["text"] for chunk in first), text)
        self.assertTrue(all(len(chunk["text"]) <= 10 for chunk in first))
        self.assertEqual([chunk["chunk_hash"] for chunk in first],
                         [chunk["chunk_hash"] for chunk in second])
        self.assertTrue(all(chunk["case_id"] == "CASE-A" for chunk in first))
        bounded_service = DocumentIngestionService(self.repo, self.embeddings,
            max_file_size=2_000_000, max_uncompressed=2_000_000, chunk_tokens=100,
            chunk_overlap=0, batch_size=10, ocr_min_chars=20, chunk_max_chars=1000)
        paragraph = "abcdefghijk " * 50
        grouped = bounded_service._chunk("CASE-A", [
            {"text": paragraph, "page": 1, "section": "HECHOS", "paragraph_start": 1, "paragraph_end": 1},
            {"text": paragraph, "page": 1, "section": "HECHOS", "paragraph_start": 2, "paragraph_end": 2},
        ])
        self.assertEqual(len(grouped), 2)
        self.assertTrue(all(len(chunk["text"]) <= 1000 for chunk in grouped))

    def test_retention_purges_document_and_chunks_by_explicit_policy(self):
        manifest, _ = self.service.ingest("CASE-A", "brief.txt", self.txt("Alpha evidence must expire."))
        with self.repo._connect() as db:
            db.execute("UPDATE case_documents SET updated_at='2000-01-01T00:00:00+00:00' WHERE document_id=?",
                       (manifest["document_id"],))
        self.assertEqual(self.repo.purge_expired(30), 1)
        self.assertEqual(self.repo.list_documents("CASE-A"), [])
        self.assertEqual(self.repo.retrieve("CASE-A"), [])

    def test_page_limit_is_enforced_before_extraction(self):
        service = DocumentIngestionService(self.repo, self.embeddings,
            max_file_size=2_000_000, max_uncompressed=2_000_000, chunk_tokens=16,
            chunk_overlap=3, batch_size=2, ocr_min_chars=20, max_pages=1)
        pdf = fitz.open()
        pdf.new_page()
        pdf.new_page()
        payload = pdf.tobytes()
        pdf.close()
        with self.assertRaisesRegex(DocumentError, "páginas"):
            service.ingest("CASE-A", "too-many.pdf", io.BytesIO(payload))

    def test_invalid_files_and_empty_text_are_rejected_safely(self):
        with self.assertRaisesRegex(DocumentError, r"\.doc.*no compatible|\.doc"):
            self.service.ingest("CASE-A", "record.doc", self.txt("binary"))
        with self.assertRaises(DocumentError):
            self.service.ingest("CASE-A", "record.pdf", self.txt("not really pdf"))
        with self.assertRaises(DocumentError):
            self.service.ingest("CASE-A", "empty.txt", self.txt("\x00\n"))
        with self.assertRaises(DocumentError):
            self.service.ingest("CASE-A", "huge.txt", io.BytesIO(b"x" * 2_000_001))


class UploadApiTests(unittest.TestCase):
    def test_document_task_reports_queued_running_and_completed(self):
        from threading import Event, Thread as RealThread
        from app.services.document_tasks import DocumentTaskManager
        entered, release = Event(), Event()
        tmp = tempfile.TemporaryDirectory()
        staging = Path(tmp.name) / "staging"
        staging.mkdir()
        source = staging / "brief.txt"
        source.write_text("ready", encoding="utf-8")
        class SlowIngestion:
            def ingest(self, case_id, filename, stream, progress, document_id, **kwargs):
                progress("extracting", 1)
                entered.set()
                if not release.wait(3):
                    raise TimeoutError("test release timeout")
                return ({"document_id":document_id, "case_id":case_id, "filename":filename,
                         "status":"ready", "page_count":None, "chunk_count":1,
                         "ocr_required":False, "warnings":[]}, False)
        class DeferredThread:
            def __init__(self, target, args, **kwargs): self.target, self.args = target, args
            def start(self): pass
        manager = DocumentTaskManager()
        try:
            with patch("app.services.document_tasks.Thread", DeferredThread):
                task_id = manager.start("CASE-A", [(source,"brief.txt")], SlowIngestion(), str(staging))
            self.assertEqual(manager.status("CASE-A", task_id)["status"], "queued")
            worker = RealThread(target=manager._run, args=(task_id, "CASE-A",
                [(source,"brief.txt",None,manager.status("CASE-A",task_id)["documents"][0]["document_id"])],
                SlowIngestion(), str(staging), None))
            worker.start()
            self.assertTrue(entered.wait(2))
            running = manager.status("CASE-A", task_id)
            self.assertEqual(running["status"], "running")
            self.assertEqual(running["stage"], "extracting")
            release.set()
            worker.join(2)
            self.assertEqual(manager.status("CASE-A", task_id)["status"], "completed")
            self.assertFalse(staging.exists())
        finally:
            release.set()
            tmp.cleanup()

    def test_document_task_partial_failure_case_scope_and_staging_cleanup(self):
        import time
        from app.services.document_tasks import DocumentTaskManager
        from app.services.case_corpus import DocumentError
        tmp = tempfile.TemporaryDirectory()
        root = Path(tmp.name)
        staging = root / "staging"
        staging.mkdir()
        first, second = staging / "first.txt", staging / "second.txt"
        first.write_text("valid", encoding="utf-8")
        second.write_text("bad", encoding="utf-8")
        class PartialIngestion:
            def ingest(self, case_id, filename, stream, **kwargs):
                self.assert_case = case_id
                kwargs["progress"]("extracting", 1)
                if filename == "bad.txt":
                    raise DocumentError("parse_failed", "No se pudo procesar el archivo.")
                return ({"document_id":"DOC-A", "case_id":case_id, "filename":filename,
                         "status":"ready", "page_count":None, "chunk_count":1,
                         "ocr_required":False, "warnings":[]}, False)
        manager = DocumentTaskManager()
        task_id = manager.start("CASE-A", [(first,"good.txt"),(second,"bad.txt")],
                                PartialIngestion(), str(staging))
        deadline = time.monotonic() + 3
        status = None
        while time.monotonic() < deadline:
            status = manager.status("CASE-A", task_id)
            if status and status["status"] in {"completed", "failed"}:
                break
            time.sleep(0.01)
        self.assertEqual(status["status"], "completed")
        self.assertEqual([item["status"] for item in status["documents"]], ["ready", "failed"])
        self.assertEqual(status["documents"][1]["case_id"], "CASE-A")
        self.assertTrue(status["documents"][1]["document_id"])
        self.assertIsNone(status["documents"][1]["page_count"])
        self.assertEqual(status["documents"][1]["chunk_count"], 0)
        self.assertIsNone(manager.status("CASE-B", task_id))
        self.assertFalse(staging.exists())
        tmp.cleanup()

    def test_task_pipeline_forwards_explicit_case_and_document_ids(self):
        import time
        from app.services.novacourt_pipeline_service import NovaCourtPipelineService
        class FakeCaseService:
            def __init__(self): self.calls = []
            def analyze_case(self, text, **kwargs):
                self.calls.append((text, kwargs["case_id"], kwargs["document_ids"]))
                kwargs["progress_callback"]("intake", "completed")
                return {"success":True, "case_id":kwargs["case_id"], "case":text}
        service = FakeCaseService()
        pipeline = NovaCourtPipelineService(service, Mock(), Mock())
        task_id = pipeline.start("Case narrative", case_id="CASE-A", tool="case", document_ids=["DOC-A"])
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            status = pipeline.status(task_id)
            if status and status["status"] in {"completed", "failed", "cancelled"}:
                break
            time.sleep(0.01)
        self.assertEqual(status["status"], "completed")
        self.assertEqual(service.calls, [("Case narrative", "CASE-A", ["DOC-A"])])

    def test_novacourt_continuation_keeps_the_explicit_corpus_identity(self):
        from app.services.novacourt_pipeline_service import NovaCourtPipelineService
        from app.services.source_contracts import normalize_case_source
        source = normalize_case_source({"case_id":"CASE-A", "document_id":"DOC-A", "chunk_id":"CHUNK-A",
                                       "text":"Same-case evidence"})
        class FakeCaseService:
            def __init__(self): self.received = None
            def analyze_case(self, text, **kwargs):
                self.received = (kwargs["case_id"], list(kwargs["document_ids"]))
                kwargs["progress_callback"]("intake", "completed")
                return {"success":True, "case_id":kwargs["case_id"], "case":text, "sources":[source],
                        "citations":[{"citation_id":"CIT-123456789012345678901234",
                                       "source_id":source["source_id"], "label":"[1]"}]}
        case = FakeCaseService()
        graph = Mock()
        graph.build_for_case.return_value = {"status":"ready", "nodes":[], "edges":[]}
        simulation = Mock()
        simulation.simulate_for_case.return_value = {"status":"ready", "participants":[], "projection":{}}
        pipeline = NovaCourtPipelineService(case, graph, simulation)
        task_id = pipeline.start("Continue this case", case_id="CASE-A", tool="court", document_ids=["DOC-A"])
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            status = pipeline.status(task_id)
            if status and status["status"] in {"completed", "failed", "cancelled"}:
                break
            time.sleep(0.01)
        self.assertEqual(status["status"], "completed")
        self.assertEqual(case.received, ("CASE-A", ["DOC-A"]))
        self.assertEqual(status["final_result"]["graph"]["status"], "ready")
        self.assertEqual(status["final_result"]["simulation"]["status"], "ready")
        self.assertEqual(status["final_result"]["sources"][0]["case_id"], "CASE-A")
        self.assertEqual(status["final_result"]["citations"][0]["source_id"], source["source_id"])

    def test_case_service_uses_only_explicit_case_corpus_with_source_references(self):
        from tests.test_stabilization import case_service
        tmp = tempfile.TemporaryDirectory()
        repository = CaseCorpusRepository(str(Path(tmp.name) / "case-corpus.sqlite3"))
        ingestion = DocumentIngestionService(repository, FakeEmbeddings(),
            max_file_size=100_000, max_uncompressed=100_000, chunk_tokens=64,
            chunk_overlap=5, batch_size=8, ocr_min_chars=20)
        manifest, _ = ingestion.ingest("CASE-A", "evidence.txt", io.BytesIO(
            b"Alpha evidence establishes the date of service and the response was late."))
        service = case_service()
        service.case_corpus_repository = repository
        try:
            with patch("app.services.embedding_service.EmbeddingService", return_value=self.fake_embeddings()):
                result = service.analyze_case("Analyze the event", case_id="CASE-A",
                                              document_ids=[manifest["document_id"]])
            analyzed_text = service.case_analyzer.analyze_case.call_args.args[0]
            self.assertIn("Alpha evidence establishes", analyzed_text)
            self.assertEqual(result["metadata"]["case_id"], "CASE-A")
            self.assertEqual(result["research"]["case_sources"][0]["document_id"],
                             manifest["document_id"])
            private_source = next(source for source in result["sources"] if source["source_scope"] == "case")
            self.assertEqual(private_source["case_id"], "CASE-A")
            self.assertEqual(private_source["document_id"], manifest["document_id"])
            self.assertIn("Alpha evidence establishes", private_source["excerpt"])
            with self.assertRaises(DocumentError):
                service.case_corpus_repository.retrieve("CASE-B", [manifest["document_id"]])
        finally:
            tmp.cleanup()

    @staticmethod
    def fake_embeddings():
        return FakeEmbeddings()

    def test_upload_list_delete_and_source_are_case_scoped(self):
        from app import create_app
        from app.config import Config
        tmp = tempfile.TemporaryDirectory()
        repo = CaseCorpusRepository(str(Path(tmp.name) / "api.sqlite3"))
        fake = FakeEmbeddings()
        class TestConfig(Config):
            TESTING = True
            RATE_LIMIT_ENABLED = False
            CASE_CORPUS_REPOSITORY_FACTORY = staticmethod(lambda: repo)
            DOCUMENT_EMBEDDING_FACTORY = staticmethod(lambda: fake)
        try:
            client = create_app(TestConfig).test_client()
            rejected = client.post("/api/cases/CASE-A/documents", data={
                "files": [(io.BytesIO(b"A"), "one.txt"), (io.BytesIO(b"B"), "two.txt")],
                "replaces_document_id":"DOC-OLD"}, content_type="multipart/form-data")
            self.assertEqual(rejected.status_code, 400)
            response = client.post("/api/cases/CASE-A/documents", data={
                "files": [(io.BytesIO(b"Alpha evidence from a valid TXT upload."), "../brief.txt"),
                          (io.BytesIO(b"old binary format"), "legacy.doc")]},
                content_type="multipart/form-data")
            self.assertEqual(response.status_code, 202)
            task_id = response.json["task_id"]
            self.assertEqual(response.json["total_documents"], 2)
            self.assertEqual(client.get(f"/api/cases/CASE-B/document-tasks/{task_id}").status_code, 404)
            deadline = time.monotonic() + 3
            task = None
            while time.monotonic() < deadline:
                task = client.get(f"/api/cases/CASE-A/document-tasks/{task_id}").json
                if task["status"] in {"completed", "failed"}:
                    break
                time.sleep(0.01)
            self.assertEqual(task["status"], "completed")
            manifest = task["documents"][0]["document"]
            self.assertEqual(task["documents"][1]["status"], "failed")
            self.assertEqual(task["documents"][1]["error"]["code"], "unsupported_format")
            chunk = repo.retrieve("CASE-A", [manifest["document_id"]])[0]
            source = client.get(f"/api/cases/CASE-A/documents/{manifest['document_id']}/chunks/{chunk['chunk_id']}")
            self.assertEqual(source.json["chunk"]["case_id"], "CASE-A")
            self.assertEqual(client.get("/api/cases/CASE-B/documents").json["documents"], [])
            self.assertEqual(client.get(f"/api/cases/CASE-B/documents/{manifest['document_id']}/chunks/{chunk['chunk_id']}").status_code, 404)
            self.assertEqual(client.delete(f"/api/cases/CASE-B/documents/{manifest['document_id']}").status_code, 404)
            self.assertEqual(client.delete(f"/api/cases/CASE-A/documents/{manifest['document_id']}").status_code, 200)
        finally:
            tmp.cleanup()

    def test_oversized_http_upload_has_safe_specific_error(self):
        from app import create_app
        from app.config import Config
        class TinyConfig(Config):
            TESTING = True
            RATE_LIMIT_ENABLED = False
            MAX_CONTENT_LENGTH = 32
        response = create_app(TinyConfig).test_client().post("/api/cases/CASE-A/documents", data={
            "file":(io.BytesIO(b"x" * 100), "brief.txt")}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json["error"]["code"], "FILE_TOO_LARGE")
        self.assertNotIn("traceback", response.json)


class RuntimeHardeningTests(unittest.TestCase):
    def test_staging_cleanup_removes_only_old_owned_directories(self):
        import os
        from app.services.document_tasks import cleanup_abandoned_staging
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            old = root / "nova-document-task-old123"
            recent = root / "nova-document-task-recent123"
            unrelated = root / "other-service-task-old123"
            for directory in (old, recent, unrelated):
                directory.mkdir()
            (old / "payload.tmp").write_text("temporary", encoding="utf-8")
            os.utime(old, (1, 1))
            os.utime(recent, (999, 999))

            with patch("app.services.document_tasks.tempfile.gettempdir", return_value=temporary):
                removed = cleanup_abandoned_staging(max_age_seconds=100, now=1000)

            self.assertEqual(removed, 1)
            self.assertFalse(old.exists())
            self.assertTrue(recent.exists())
            self.assertTrue(unrelated.exists())

    def test_staging_cleanup_skips_symlinks_and_cleanup_errors(self):
        import os
        from app.services import document_tasks
        with tempfile.TemporaryDirectory() as temporary, tempfile.TemporaryDirectory() as target:
            root = Path(temporary)
            old = root / "nova-document-task-old123"
            old.mkdir()
            os.utime(old, (1, 1))
            link = root / "nova-document-task-link123"
            try:
                link.symlink_to(target, target_is_directory=True)
            except (OSError, NotImplementedError):
                self.skipTest("Directory symlinks are unavailable on this platform")
            with patch("app.services.document_tasks.tempfile.gettempdir", return_value=temporary):
                with patch.object(document_tasks.shutil, "rmtree", side_effect=OSError("locked")):
                    manager = document_tasks.DocumentTaskManager(staging_max_age_seconds=100)
                    self.assertIsNotNone(manager)
            self.assertTrue(old.exists())
            self.assertTrue(Path(target).exists())

    def test_legacy_pdf_route_uses_pypdf_and_tolerates_empty_page_text(self):
        from types import SimpleNamespace
        from app.api.simulation import extraer_texto_documento
        upload = SimpleNamespace(filename="brief.pdf")
        with patch("app.api.simulation.PdfReader", return_value=SimpleNamespace(pages=[
            SimpleNamespace(extract_text=lambda: "first page"),
            SimpleNamespace(extract_text=lambda: None),
            SimpleNamespace(extract_text=lambda: "third page"),
        ])) as reader:
            self.assertEqual(extraer_texto_documento(upload), "first page\n\nthird page")
        reader.assert_called_once_with(upload)
