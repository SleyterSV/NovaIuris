"""Case-isolated document ingestion and retrieval backed by local SQLite."""
from __future__ import annotations

import codecs
import hashlib
import heapq
import json
import math
import os
import re
import sqlite3
import tempfile
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, BinaryIO


class DocumentError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def safe_filename(filename: str | None) -> str:
    name = (filename or "").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[^\w .()\[\]-]", "_", name, flags=re.UNICODE).strip(" .")
    if not name or name in {".", ".."}:
        raise DocumentError("invalid_filename", "El nombre de archivo no es válido.")
    return name[:180]


def _word_spans(text: str) -> list[tuple[int, int]]:
    return [(m.start(), m.end()) for m in re.finditer(r"\w+|[^\w\s]", text, re.UNICODE)]


def _section(text: str, current: str | None) -> str | None:
    headings = ("FUNDAMENTOS DE HECHO", "FUNDAMENTOS DE DERECHO", "PETITORIO",
                "MEDIOS PROBATORIOS", "ANEXOS", "CONSIDERANDOS", "RESUELVE")
    normalized = re.sub(r"[^A-ZÁÉÍÓÚÜÑ ]", "", text.upper()).strip()
    for heading in headings:
        if normalized.startswith(heading) and len(normalized) <= len(heading) + 80:
            return heading
    return current


def _decode_txt(path: Path) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
        parts = []
        try:
            with path.open("rb") as source:
                while block := source.read(1024 * 1024):
                    parts.append(decoder.decode(block))
                parts.append(decoder.decode(b"", final=True))
            text = "".join(parts)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocumentError("parse_failed", "No se pudo decodificar el archivo de texto.")
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(re.sub(r"[\t ]+", " ", line).strip() for line in text.split("\n")).strip()


def _docx_units(path: Path, max_uncompressed: int) -> list[dict[str, Any]]:
    from docx import Document
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > max_uncompressed:
                raise DocumentError("file_too_large", "El DOCX supera el límite de extracción segura.")
            if "word/document.xml" not in archive.namelist():
                raise DocumentError("mime_mismatch", "El archivo no contiene un documento DOCX válido.")
    except zipfile.BadZipFile as error:
        raise DocumentError("mime_mismatch", "El archivo no contiene un documento DOCX válido.") from error

    document = Document(str(path))
    units, section, paragraph_no = [], None, 0
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            paragraph_no += 1
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                section = text[:180]
            section = _section(text, section)
            units.append({"text": text, "page": None, "section": section,
                          "paragraph_start": paragraph_no, "paragraph_end": paragraph_no})
        elif tag == "tbl":
            from docx.table import Table
            table = Table(child, document)
            for row_no, row in enumerate(table.rows, start=1):
                cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                text = " | ".join(value for value in cells if value)
                if text:
                    paragraph_no += 1
                    units.append({"text": text, "page": None, "section": section,
                                  "paragraph_start": paragraph_no, "paragraph_end": paragraph_no,
                                  "table_row": row_no})
    return units


class CaseCorpusRepository:
    """Persistent private text/vector storage; every read requires case_id."""
    def __init__(self, database_path: str):
        self.database_path = database_path
        Path(database_path).parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS case_documents (
                    document_id TEXT PRIMARY KEY, case_id TEXT NOT NULL,
                    sha256 TEXT NOT NULL, version INTEGER NOT NULL, status TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    active INTEGER NOT NULL DEFAULT 1, manifest_json TEXT NOT NULL,
                    UNIQUE(case_id, sha256), UNIQUE(document_id, case_id));
                CREATE INDEX IF NOT EXISTS ix_case_documents_case ON case_documents(case_id, active);
                CREATE TABLE IF NOT EXISTS case_chunks (
                    chunk_id TEXT PRIMARY KEY, case_id TEXT NOT NULL, document_id TEXT NOT NULL,
                    chunk_hash TEXT NOT NULL, page_start INTEGER, page_end INTEGER,
                    section TEXT, source_json TEXT NOT NULL, text TEXT NOT NULL,
                    embedding_json TEXT NOT NULL,
                    FOREIGN KEY(document_id, case_id) REFERENCES case_documents(document_id, case_id));
                CREATE INDEX IF NOT EXISTS ix_case_chunks_case_doc ON case_chunks(case_id, document_id);
            """)

    @contextmanager
    def _connect(self):
        db = sqlite3.connect(self.database_path, timeout=30)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        try:
            yield db
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def find_hash(self, case_id: str, digest: str):
        if not case_id:
            raise DocumentError("case_id_required", "Se requiere case_id.")
        with self._connect() as db:
            row = db.execute("SELECT manifest_json FROM case_documents WHERE case_id=? AND sha256=?", (case_id, digest)).fetchone()
        return json.loads(row[0]) if row else None

    def save(self, manifest: dict, chunks: list[dict], replaces_document_id: str | None = None):
        case_id = manifest["case_id"]
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            existing = db.execute("SELECT manifest_json FROM case_documents WHERE case_id=? AND sha256=?",
                                  (case_id, manifest["sha256"])).fetchone()
            if existing:
                old = json.loads(existing[0])
                return old, True
            if replaces_document_id:
                old = db.execute("SELECT manifest_json FROM case_documents WHERE document_id=? AND case_id=? AND active=1",
                                 (replaces_document_id, case_id)).fetchone()
                if old is None:
                    raise DocumentError("document_not_found", "El documento que se desea reemplazar no pertenece a este caso.")
                version = json.loads(old[0]).get("document_version", 1) + 1
                db.execute("DELETE FROM case_chunks WHERE case_id=? AND document_id=?",
                           (case_id, replaces_document_id))
                db.execute("DELETE FROM case_documents WHERE document_id=? AND case_id=?",
                           (replaces_document_id, case_id))
                manifest["document_version"] = version
            db.execute("INSERT INTO case_documents(document_id,case_id,sha256,version,status,updated_at,manifest_json) VALUES(?,?,?,?,?,?,?)",
                       (manifest["document_id"], case_id, manifest["sha256"], manifest["document_version"], manifest["status"],
                        manifest["updated_at"],
                        json.dumps(manifest, ensure_ascii=False)))
            db.executemany("""INSERT INTO case_chunks(chunk_id,case_id,document_id,chunk_hash,page_start,page_end,
                          section,source_json,text,embedding_json) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                [(c["chunk_id"], case_id, manifest["document_id"], c["chunk_hash"], c.get("page_start"),
                  c.get("page_end"), c.get("section"), json.dumps(c["source_reference"], ensure_ascii=False),
                  c["text"], json.dumps(c["embedding"])) for c in chunks])
            return manifest, False

    def list_documents(self, case_id: str):
        if not case_id:
            raise DocumentError("case_id_required", "Se requiere case_id.")
        with self._connect() as db:
            rows = db.execute("SELECT manifest_json FROM case_documents WHERE case_id=? AND active=1 ORDER BY rowid", (case_id,)).fetchall()
        return [json.loads(row[0]) for row in rows]

    def case_size_and_count(self, case_id: str):
        if not case_id:
            raise DocumentError("case_id_required", "Se requiere case_id.")
        with self._connect() as db:
            row = db.execute("SELECT COALESCE(SUM(json_extract(manifest_json,'$.size_bytes')),0), COUNT(*) FROM case_documents WHERE case_id=? AND active=1", (case_id,)).fetchone()
        return int(row[0]), int(row[1])

    def delete_document(self, case_id: str, document_id: str):
        if not case_id:
            raise DocumentError("case_id_required", "Se requiere case_id.")
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            exists = db.execute("SELECT 1 FROM case_documents WHERE case_id=? AND document_id=?", (case_id, document_id)).fetchone()
            if not exists:
                return False
            db.execute("DELETE FROM case_chunks WHERE case_id=? AND document_id=?", (case_id, document_id))
            db.execute("DELETE FROM case_documents WHERE case_id=? AND document_id=?", (case_id, document_id))
            return True

    def get_chunk(self, case_id: str, document_id: str, chunk_id: str):
        if not case_id:
            raise DocumentError("case_id_required", "La consulta privada requiere case_id.")
        with self._connect() as db:
            row = db.execute("SELECT source_json,text FROM case_chunks WHERE case_id=? AND document_id=? AND chunk_id=?",
                             (case_id, document_id, chunk_id)).fetchone()
        if row is None:
            return None
        return {"case_id": case_id, "document_id": document_id, "chunk_id": chunk_id,
                "source_reference": json.loads(row[0]), "text": row[1]}

    def validate_documents(self, case_id: str, document_ids: list[str] | None = None):
        if not case_id:
            raise DocumentError("case_id_required", "La búsqueda privada requiere case_id.")
        with self._connect() as db:
            active = {row[0] for row in db.execute(
                "SELECT document_id FROM case_documents WHERE case_id=? AND active=1 AND status='ready'", (case_id,))}
        if document_ids is not None and not set(document_ids) <= active:
            raise DocumentError("document_not_found", "Uno o más documentos no pertenecen al caso o no están listos.")
        return active if document_ids is None else set(document_ids)

    def iter_retrieve(self, case_id: str, document_ids: list[str] | None = None):
        if not case_id:
            raise DocumentError("case_id_required", "La búsqueda privada requiere case_id.")
        selected = self.validate_documents(case_id, document_ids)
        if not selected:
            return
        with self._connect() as db:
            placeholders = ",".join("?" for _ in selected)
            cursor = db.execute(f"SELECT chunk_id,case_id,document_id,page_start,page_end,section,source_json,text,embedding_json FROM case_chunks WHERE case_id=? AND document_id IN ({placeholders})",
                                [case_id, *sorted(selected)])
            while rows := cursor.fetchmany(128):
                for row in rows:
                    yield {**dict(row), "source_reference": json.loads(row["source_json"]),
                           "embedding": json.loads(row["embedding_json"])}

    def retrieve(self, case_id: str, document_ids: list[str] | None = None):
        return list(self.iter_retrieve(case_id, document_ids))

    def delete_case(self, case_id: str):
        if not case_id:
            raise DocumentError("case_id_required", "Se requiere case_id.")
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            db.execute("DELETE FROM case_chunks WHERE case_id=?", (case_id,))
            db.execute("DELETE FROM case_documents WHERE case_id=?", (case_id,))

    def purge_expired(self, retention_days: int):
        if retention_days <= 0:
            return 0
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            expired = [row[0] for row in db.execute(
                "SELECT document_id FROM case_documents WHERE updated_at < datetime('now', ?)",
                (f"-{int(retention_days)} days",))]
            if expired:
                marks = ",".join("?" for _ in expired)
                db.execute(f"DELETE FROM case_chunks WHERE document_id IN ({marks})", expired)
                db.execute(f"DELETE FROM case_documents WHERE document_id IN ({marks})", expired)
            return len(expired)


class DocumentIngestionService:
    SUPPORTED = {".pdf", ".docx", ".txt"}

    def __init__(self, repository: CaseCorpusRepository, embedding_service: Any,
                 *, max_file_size: int, max_uncompressed: int, chunk_tokens: int,
                 chunk_overlap: int, batch_size: int, ocr_min_chars: int,
                 max_pages: int = 5000, chunk_max_chars: int = 12000,
                 max_total_size: int = 200 * 1024 * 1024, max_files: int = 20):
        self.repository, self.embedding_service = repository, embedding_service
        self.max_file_size, self.max_uncompressed = max_file_size, max_uncompressed
        self.chunk_tokens, self.chunk_overlap = chunk_tokens, chunk_overlap
        self.max_pages, self.chunk_max_chars = max_pages, chunk_max_chars
        self.batch_size, self.ocr_min_chars = batch_size, ocr_min_chars
        self.max_total_size, self.max_files = max_total_size, max_files

    def ingest(self, case_id: str, filename: str, stream: BinaryIO,
               replaces_document_id: str | None = None, progress=None, document_id: str | None = None):
        if not case_id:
            raise DocumentError("case_id_required", "Se requiere case_id.")
        display_name = safe_filename(filename)
        ext = Path(display_name).suffix.lower()
        if ext == ".doc":
            raise DocumentError("unsupported_format", "Formato .doc no compatible. Convierta el documento a .docx o PDF.")
        if ext not in self.SUPPORTED:
            raise DocumentError("unsupported_format", "Formatos compatibles: PDF, DOCX y TXT.")
        digest = hashlib.sha256()
        size = 0
        too_large = False
        with tempfile.NamedTemporaryFile(prefix="nova-doc-", suffix=ext, delete=False) as temp:
            temp_path = Path(temp.name)
            try:
                while block := stream.read(1024 * 1024):
                    size += len(block)
                    if size > self.max_file_size:
                        too_large = True
                        break
                    digest.update(block)
                    temp.write(block)
            except Exception as error:
                read_error = error
            else:
                read_error = None
        if read_error is not None:
            temp_path.unlink(missing_ok=True)
            raise read_error
        if too_large:
            temp_path.unlink(missing_ok=True)
            raise DocumentError("file_too_large", "El archivo supera el límite configurado.")
        try:
            if size == 0:
                raise DocumentError("empty_document", "El archivo está vacío.")
            sha = digest.hexdigest()
            existing = self.repository.find_hash(case_id, sha)
            if existing:
                return existing, True
            case_size, file_count = self.repository.case_size_and_count(case_id)
            replacing_size = 0
            if replaces_document_id:
                replaced = next((doc for doc in self.repository.list_documents(case_id)
                                 if doc["document_id"] == replaces_document_id), None)
                if replaced is None:
                    raise DocumentError("document_not_found", "El documento que se desea reemplazar no pertenece a este caso.")
                replacing_size = replaced["size_bytes"]
            if file_count - bool(replaces_document_id) >= self.max_files:
                raise DocumentError("too_many_files", "El caso alcanzó el máximo configurado de documentos.")
            if case_size - replacing_size + size > self.max_total_size:
                raise DocumentError("case_too_large", "La suma de documentos supera el límite configurado del caso.")
            if progress: progress("validating", size)
            self._validate_signature(temp_path, ext)
            if progress: progress("extracting", 0)
            units, page_count, ocr_pages = self._extract(temp_path, ext)
            if progress: progress("normalizing", len(units))
            normalized = []
            for unit in units:
                text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", unit["text"])
                text = re.sub(r"\s+", " ", text).strip()
                if text:
                    normalized.append({**unit, "text": text})
            if progress: progress("chunking", len(normalized))
            chunks = self._chunk(case_id, normalized)
            document_version = 1
            if replaces_document_id:
                replaced = next(doc for doc in self.repository.list_documents(case_id)
                               if doc["document_id"] == replaces_document_id)
                document_version = replaced.get("document_version", 1) + 1
            embeddings = []
            for offset in range(0, len(chunks), self.batch_size):
                batch = chunks[offset:offset + self.batch_size]
                embeddings.extend(self.embedding_service.generate_embeddings([c["text"] for c in batch]))
                if progress: progress("indexing", min(offset + len(batch), len(chunks)))
            for chunk, vector in zip(chunks, embeddings):
                chunk["embedding"] = vector
            warnings = [{"code": "ocr_required", "page_number": page} for page in ocr_pages]
            status = "ready" if chunks else "ocr_required" if ocr_pages else "failed"
            if status == "failed":
                raise DocumentError("empty_document", "No se pudo extraer texto útil del documento.")
            now = utc_now()
            manifest = {"document_id": document_id or str(uuid.uuid4()), "case_id": case_id,
                "filename": display_name, "display_name": Path(display_name).stem,
                "mime_type": {".pdf":"application/pdf", ".docx":"application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".txt":"text/plain"}[ext],
                "extension": ext, "size_bytes": size, "sha256": sha, "content_hash": sha,
                "document_version": document_version, "status": status, "page_count": page_count,
                "character_count": sum(len(u["text"]) for u in normalized),
                "token_count": sum(len(_word_spans(u["text"])) for u in normalized),
                "chunk_count": len(chunks), "ocr_required": bool(ocr_pages), "warnings": warnings,
                "metadata": {"document_type": None, "case_number": None, "court": None,
                    "entity": None, "date": None, "title": None, "parties": None,
                    "procedural_role": None, "source_type": "user_upload"},
                "created_at": now, "updated_at": now}
            for chunk in chunks:
                chunk["document_id"] = manifest["document_id"]
                chunk["source_reference"].update({"case_id": case_id,
                    "document_id": manifest["document_id"], "chunk_id": chunk["chunk_id"]})
            saved, duplicate = self.repository.save(manifest, chunks, replaces_document_id)
            if progress: progress("ready", len(chunks))
            return saved, duplicate
        finally:
            temp_path.unlink(missing_ok=True)

    @staticmethod
    def _validate_signature(path: Path, ext: str):
        with path.open("rb") as source: header = source.read(8)
        if ext == ".pdf" and not header.startswith(b"%PDF-"):
            raise DocumentError("mime_mismatch", "La extensión no coincide con el contenido PDF.")
        if ext == ".docx" and not header.startswith(b"PK"):
            raise DocumentError("mime_mismatch", "La extensión no coincide con el contenido DOCX.")
        if ext == ".txt" and header.startswith((b"%PDF-", b"PK")):
            raise DocumentError("mime_mismatch", "El contenido no coincide con un archivo TXT.")

    def _extract(self, path: Path, ext: str):
        if ext == ".pdf":
            import fitz
            units, ocr = [], []
            try:
                with fitz.open(str(path)) as pdf:
                    if len(pdf) > self.max_pages:
                        raise DocumentError("too_many_pages", "El PDF supera el máximo configurado de páginas.")
                    if pdf.is_encrypted:
                        raise DocumentError("parse_failed", "El PDF está protegido y no se puede leer.")
                    section = None
                    for number, page in enumerate(pdf, 1):
                        text = page.get_text("text") or ""
                        if len(re.sub(r"\s", "", text)) < self.ocr_min_chars:
                            ocr.append(number)
                        for i, paragraph in enumerate(re.split(r"\n\s*\n", text)):
                            if paragraph.strip():
                                section = _section(paragraph, section)
                                units.append({"text": paragraph, "page": number,
                                    "section": section, "paragraph_start": i + 1, "paragraph_end": i + 1})
                    return units, len(pdf), ocr
            except DocumentError: raise
            except Exception as error:
                raise DocumentError("parse_failed", "No se pudo procesar el PDF.") from error
        if ext == ".docx":
            return _docx_units(path, self.max_uncompressed), None, []
        text = _decode_txt(path)
        units, section = [], None
        for i, paragraph in enumerate(re.split(r"\n\s*\n", text), 1):
            paragraph = paragraph.strip()
            if paragraph:
                section = _section(paragraph, section)
                units.append({"text": paragraph, "page": None, "section": section,
                              "paragraph_start": i, "paragraph_end": i})
        return units, None, []

    def _chunk(self, case_id: str, units: list[dict]):
        chunks, current, current_tokens = [], [], 0
        def emit(parts):
            text = "\n\n".join(part["text"] for part in parts).strip()
            if not text: return
            page_values = [part["page"] for part in parts if part.get("page") is not None]
            refs = {"page_start": min(page_values) if page_values else None,
                    "page_end": max(page_values) if page_values else None,
                    "section": next((p.get("section") for p in reversed(parts) if p.get("section")), None),
                    "paragraph_start": parts[0].get("paragraph_start"),
                    "paragraph_end": parts[-1].get("paragraph_end"),
                    "table_row": parts[0].get("table_row")}
            digest = hashlib.sha256(re.sub(r"\s+", " ", text).strip().encode()).hexdigest()
            chunks.append({"chunk_id": str(uuid.uuid4()), "case_id": case_id,
                "document_id": None, "text": text, "chunk_hash": digest,
                "page_start": refs["page_start"], "page_end": refs["page_end"],
                "section": refs["section"], "source_reference": refs})
        for unit in units:
            spans = _word_spans(unit["text"])
            if len(spans) > self.chunk_tokens or len(unit["text"]) > self.chunk_max_chars:
                if current: emit(current); current, current_tokens = [], 0
                start = 0
                while start < len(spans):
                    end = min(start + self.chunk_tokens, len(spans))
                    while end > start + 1 and spans[end - 1][1] - spans[start][0] > self.chunk_max_chars:
                        end -= 1
                    if spans[end - 1][1] - spans[start][0] > self.chunk_max_chars:
                        begin = spans[start][0]
                        finish = spans[end - 1][1]
                        piece = unit["text"][begin:finish]
                        for offset in range(0, len(piece), self.chunk_max_chars):
                            emit([{**unit, "text": piece[offset:offset + self.chunk_max_chars]}])
                        start = end
                        continue
                    piece = unit["text"][spans[start][0]:spans[end - 1][1]]
                    emit([{**unit, "text": piece}])
                    if end == len(spans): break
                    start = max(start + 1, end - self.chunk_overlap)
                continue
            if current and (current_tokens + len(spans) > self.chunk_tokens or
                            sum(len(part["text"]) for part in current) + len(unit["text"]) + 2 * len(current) > self.chunk_max_chars or
                            current[-1].get("page") != unit.get("page") or
                            current[-1].get("section") != unit.get("section")):
                emit(current)
                overlap = []
                last = current[-1]
                if (self.chunk_overlap and last.get("section") == unit.get("section")
                        and last.get("page") == unit.get("page")):
                    spans = _word_spans(last["text"])
                    if spans:
                        first = max(0, len(spans) - self.chunk_overlap)
                        overlap_text = last["text"][spans[first][0]:]
                        overlap = [{**last, "text": overlap_text}]
                if sum(len(_word_spans(p["text"])) for p in overlap) + len(spans) > self.chunk_tokens:
                    overlap = []
                current, current_tokens = overlap[:], sum(len(_word_spans(p["text"])) for p in overlap)
            current.append(unit); current_tokens += len(spans)
        emit(current)
        return chunks


class CaseContextService:
    """Retrieves ranked, traceable passages from exactly one case corpus."""
    def __init__(self, repository: CaseCorpusRepository, embedding_service: Any):
        self.repository, self.embedding_service = repository, embedding_service

    def relevant_context(self, case_id: str, query: str, *, document_ids=None, top_k=6, filters=None,
                         cancellation_token=None):
        if not case_id: raise DocumentError("case_id_required", "La búsqueda privada requiere case_id.")
        if not query or not query.strip(): return []
        if not self.repository.validate_documents(case_id, document_ids):
            return []
        query_embedding = self.embedding_service.generate_embedding(query, cancellation_token=cancellation_token)
        candidates = self.repository.iter_retrieve(case_id, document_ids)
        if filters and filters.get("section"):
            candidates = (c for c in candidates if c.get("section") == filters["section"])
        def cosine(left, right):
            if not left or not right or len(left) != len(right): return -1.0
            denom = math.sqrt(sum(v*v for v in left) * sum(v*v for v in right))
            return sum(a*b for a,b in zip(left,right)) / denom if denom else -1.0
        limit = max(1, min(int(top_k), 20))
        ranked = heapq.nlargest(limit, ((cosine(query_embedding, c["embedding"]), c) for c in candidates),
                                key=lambda item: item[0])
        return [{"text": c["text"], "source_reference": c["source_reference"],
                 "document_id": c["document_id"], "chunk_id": c["chunk_id"],
                 "similarity": score} for score, c in ranked]
