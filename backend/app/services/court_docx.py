"""Editable A4 DOCX built from NovaCourt's persisted report projection."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .case_docx import _runs
from .court_report_content import court_report_blocks, court_report_content


def _page_field(paragraph):
    run = paragraph.add_run()
    begin, instruction, separate, end = (OxmlElement('w:fldChar'), OxmlElement('w:instrText'),
                                         OxmlElement('w:fldChar'), OxmlElement('w:fldChar'))
    begin.set(qn('w:fldCharType'), 'begin')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    separate.set(qn('w:fldCharType'), 'separate')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend((begin, instruction, separate, end))


def render_court_docx(result):
    content = court_report_content(result)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.4)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Aptos', Pt(11)
    normal.font.color.rgb = RGBColor(38, 56, 75)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(9)
    for name, size in (('Title', 20), ('Heading 1', 13), ('Heading 2', 11)):
        style = document.styles[name]
        style.font.name, style.font.size = 'Aptos Display', Pt(size)
        style.font.color.rgb = RGBColor(23, 55, 94)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(16), Pt(7)

    eyebrow = document.add_paragraph('MYKE · LEGAL INTELLIGENCE')
    eyebrow.runs[0].font.size = Pt(8)
    eyebrow.runs[0].font.color.rgb = RGBColor(151, 119, 75)
    document.add_paragraph(content['title'], style='Title')
    document.add_paragraph('NovaCourt')
    if content['case_title']:
        document.add_heading(content['case_title'], level=2)
    disclaimer = document.add_paragraph('Simulación jurídica. No constituye una resolución judicial real ni sustituye la evaluación profesional.')
    disclaimer.runs[0].font.size = Pt(9)
    document.add_paragraph(f"Generado: {content['date']}")

    for item in content['sections']:
        document.add_heading(item['title'], level=1)
        for kind, value in court_report_blocks(item['content']):
            if kind == 'heading':
                document.add_heading(value, level=2)
            elif kind == 'list':
                _runs(document.add_paragraph(style='List Bullet'), value)
            elif kind == 'numbered':
                _runs(document.add_paragraph(style='List Number'), value)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Cm(.35)
                _runs(paragraph, value)

    if content['references']:
        document.add_heading('Fuentes utilizadas', level=1)
        for group in ('DOCUMENTOS DEL CASO', 'FUENTES JURÍDICAS'):
            refs = [ref for ref in content['references'] if ref['group'] == group]
            if not refs:
                continue
            document.add_heading(group, level=2)
            for reference in refs:
                paragraph = document.add_paragraph()
                paragraph.add_run(reference['label'] + ' ').bold = True
                paragraph.add_run(reference['title'])
                if reference['detail']:
                    paragraph.add_run(' · ' + reference['detail'])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Generado por MYKE · NovaCourt  ·  ').font.size = Pt(8)
    _page_field(footer)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
