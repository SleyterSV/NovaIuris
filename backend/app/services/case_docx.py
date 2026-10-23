"""Editable A4 export of the same persisted NovaCase report used for PDF."""
from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from .case_report_content import report_blocks, report_content


def _runs(paragraph, text):
    for index, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1


def render_case_docx(result):
    content = report_content(result)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.4)
    section.left_margin = section.right_margin = Cm(2.5)
    styles = document.styles
    normal = styles['Normal']
    normal.font.name, normal.font.size = 'Aptos', Pt(11)
    normal.font.color.rgb = RGBColor(38, 56, 75)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(9)
    for name, size in (('Title', 20), ('Heading 1', 13), ('Heading 2', 11)):
        style = styles[name]
        style.font.name, style.font.size = 'Aptos Display', Pt(size)
        style.font.color.rgb = RGBColor(23, 55, 94)
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(7)

    eyebrow = document.add_paragraph('NOVACASE · INFORME JURÍDICO')
    eyebrow.runs[0].font.size = Pt(8)
    eyebrow.runs[0].font.color.rgb = RGBColor(151, 119, 75)
    document.add_paragraph(content['title'], style='Title')
    document.add_paragraph('Documento de análisis preparado para revisión profesional.')
    if content['date']:
        document.add_paragraph(f"Generado: {content['date']}")

    for item in content['sections']:
        document.add_heading(item['title'], level=1)
        for kind, value in report_blocks(item['content'], item['title']):
            if kind == 'heading':
                document.add_heading(value, level=2)
            elif kind == 'list':
                _runs(document.add_paragraph(style='List Bullet'), value)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Cm(.35)
                _runs(paragraph, value)

    if content['references']:
        document.add_heading('Fuentes utilizadas', level=1)
        document.add_paragraph('Documentos del caso y fuentes jurídicas citadas en el informe.')
        for reference in content['references']:
            paragraph = document.add_paragraph()
            paragraph.add_run(reference['label'] + ' ').bold = True
            paragraph.add_run(reference['title'])
            if reference['detail']:
                paragraph.add_run(' · ' + reference['detail'])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Generado por MYKE').font.size = Pt(8)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
